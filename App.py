import google.generativeai as genai
import json
from PIL import Image
import streamlit as st
import io
import math
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from datetime import date

# ----------------- NUMBER TO WORDS CONVERTER -----------------
def amount_to_words(n):
    """Converts a number to Nigerian Naira English words."""
    if n == 0: return "Zero Naira"
    
    units = ["", "One", "Two", "Three", "Four", "Five", "Six", "Seven", "Eight", "Nine", "Ten", "Eleven", "Twelve", "Thirteen", "Fourteen", "Fifteen", "Sixteen", "Seventeen", "Eighteen", "Nineteen"]
    tens = ["", "", "Twenty", "Thirty", "Forty", "Fifty", "Sixty", "Seventy", "Eighty", "Ninety"]
    
    def convert(num):
        if num < 20: return units[num]
        if num < 100: return tens[num // 10] + ("-" + units[num % 10] if num % 10 != 0 else "")
        if num < 1000: return units[num // 100] + " Hundred" + (" and " + convert(num % 100) if num % 100 != 0 else "")
        
        for idx, word in enumerate(["Thousand", "Million", "Billion"], 1):
            if num < 1000 ** (idx + 1):
                return convert(num // (1000 ** idx)) + " " + word + (", " + convert(num % (1000 ** idx)) if num % (1000 ** idx) != 0 else "")

    n_int = int(math.floor(n))
    kobo = int(round((n - n_int) * 100))
    
    words = convert(n_int) + " Naira"
    if kobo > 0:
        words += f" and {convert(kobo)} Kobo"
        
    return words

# ----------------- BRANDING & CSS -----------------
st.set_page_config(page_title="Agege L&D Memo App", page_icon="📝", layout="centered")

st.markdown("""
<style>
    .stButton>button {
        background-color: #005a2b !important;
        color: white !important;
        border-radius: 8px !important;
        font-weight: bold !important;
        width: 100% !important;
        padding: 12px !important;
        border: none !important;
    }
    .stButton>button:hover { background-color: #003a1b !important; }
    .logo-container {
        display: flex;
        align-items: center;
        gap: 15px;
        background: linear-gradient(90deg, #005a2b, #00401d);
        padding: 15px;
        border-radius: 12px;
        color: white;
        margin-bottom: 25px;
    }
</style>
""", unsafe_allow_html=True)

# ----------------- SECURE PASS LOCK -----------------
st.markdown("""
<div class="logo-container">
    <div style="font-size: 32px;">🇳🇬</div>
    <div>
        <h2 style='margin:0; font-size: 20px; color:white;'>AGEGE LOCAL GOVERNMENT</h2>
        <p style='margin:0; font-size:12px; opacity:0.9;'>Learning & Development Unit (Admin & HR)</p>
    </div>
</div>
""", unsafe_allow_html=True)

user_password = st.text_input("🔑 Enter Unit Access Password", type="password")

if user_password != st.secrets["APP_PASSWORD"]:
    st.warning("🔒 Please enter the correct password to unlock the Memo Generator.")
    st.stop()

# ----------------- ESTACODE ENGINE -----------------
def calculate_allowance(grade_level, days, travel_type):
    estacode_rates = {
        "GL 15-17": {
            "Air": {"Accom": 165000, "Airfare": 440000, "Transp": 36000, "Incidental": 27000},
            "Road": {"Accom": 165000, "Airfare": 0, "Transp": 54000, "Incidental": 27000},
            "Within Lagos": {"Accom": 0, "Airfare": 0, "Transp": 15600, "Incidental": 27000}
        },
        "GL 13-14": {
            "Air": {"Accom": 124875, "Airfare": 440000, "Transp": 32400, "Incidental": 20250},
            "Road": {"Accom": 124875, "Airfare": 0, "Transp": 43200, "Incidental": 20250},
            "Within Lagos": {"Accom": 0, "Airfare": 0, "Transp": 9750, "Incidental": 20250}
        },
        "GL 10-12": {
            "Air": {"Accom": 82800, "Airfare": 440000, "Transp": 21600, "Incidental": 15187.50},
            "Road": {"Accom": 82800, "Airfare": 0, "Transp": 36000, "Incidental": 15187.50},
            "Within Lagos": {"Accom": 0, "Airfare": 0, "Transp": 5850, "Incidental": 15187.50}
        },
        "GL 07-09": {
            "Air": {"Accom": 62100, "Airfare": 440000, "Transp": 14400, "Incidental": 10125},
            "Road": {"Accom": 62100, "Airfare": 0, "Transp": 27000, "Incidental": 10125},
            "Within Lagos": {"Accom": 0, "Airfare": 0, "Transp": 3900, "Incidental": 10125}
        }
    }
    rates = estacode_rates.get(grade_level, estacode_rates["GL 10-12"])
    selected = rates[travel_type]
    
    accom_total = selected["Accom"] * days
    
    if travel_type == "Within Lagos":
        trans_total = selected["Transp"] * days
        inc_total = selected["Incidental"]
        sub_total = trans_total + inc_total
    else:
        trans_total = selected["Transp"]
        inc_total = selected["Incidental"]
        sub_total = selected["Airfare"] + accom_total + trans_total + inc_total
        
    return {
        "Accom_Daily": selected["Accom"],
        "Accom_Total": accom_total,
        "Airfare": selected["Airfare"],
        "Transp": trans_total,
        "Incidental": inc_total,
        "Sub_Total": sub_total
    }

# ----------------- DOCUMENT GENERATOR -----------------
def generate_memo_docx(subject, memo_date, opening_phrase, participants, fee_per_person, protocol):
    valid_participants = [p for p in participants if p.get('Name').strip() != ""]
    if not valid_participants:
        return None

    participants_subtotal = 0
    calculated_data = []

    for p in valid_participants:
        costs = calculate_allowance(p['GL'], p['Days'], p['Travel_Mode'])
        participants_subtotal += costs['Sub_Total']
        calculated_data.append({"info": p, "costs": costs})

    total_participatory_fee = fee_per_person * len(valid_participants)
    grand_total = participants_subtotal + total_participatory_fee + protocol
    amount_in_words = amount_to_words(grand_total)

    doc = Document()
    header = doc.add_paragraph()
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header.add_run("AGEGE LOCAL GOVERNMENT\nMEMORANDUM").bold = True
    header.runs[0].font.size = Pt(14)

    doc.add_paragraph(f"From: Learning & Development Unit\nAdmin & HR Department\nAgege Local Government\nRef: AGLG/GEN/177/007\n\nTo: The Executive Chairman\nThru: The Council Manager\nThru: The Human Resource Officer\nDate: {memo_date}")
    doc.add_paragraph().add_run(f"RE: {subject.upper()}").bold = True
    
    doc.add_paragraph(f"{opening_phrase} I am directed to respectfully request approval for the release of funds in respect of the above-referenced programme.")
    doc.add_paragraph(f"Accordingly, the sum of {amount_in_words} (₦{grand_total:,.2f}) only is required for sponsorship of the participants.\n\nThe breakdown is as follows:")

    for index, data in enumerate(calculated_data, start=1):
        p = data["info"]
        costs = data["costs"]
        
        doc.add_paragraph().add_run(f"{index}. {p['Name']} | {p['Title']} [{p['GL']}]").bold = True
        
        if costs['Airfare'] > 0:
            doc.add_paragraph(f"(a) Airfare: ₦{costs['Airfare']:,.2f}")
        if costs['Accom_Total'] > 0:
            doc.add_paragraph(f"(b) Accommodation and feeding: ₦{costs['Accom_Total']:,.2f}")
        doc.add_paragraph(f"(c) Transportation: ₦{costs['Transp']:,.2f}")
        doc.add_paragraph(f"(d) Incidental: ₦{costs['Incidental']:,.2f}")
        doc.add_paragraph(f"SUB-TOTAL: ₦{costs['Sub_Total']:,.2f}").runs[0].bold = True

    doc.add_paragraph(f"Participants Sub-total: ₦{participants_subtotal:,.2f}")
    if total_participatory_fee > 0: 
        doc.add_paragraph(f"Participatory Fees (₦{fee_per_person:,.2f} x {len(valid_participants)}): ₦{total_participatory_fee:,.2f}")
    if protocol > 0: 
        doc.add_paragraph(f"Protocol: ₦{protocol:,.2f}")
    
    doc.add_paragraph(f"GRAND TOTAL: ₦{grand_total:,.2f}").runs[0].bold = True
    doc.add_paragraph(f"Your kind approval is hereby sought for the sum of ₦{grand_total:,.2f}.\n\nThank you.\n\nOLATOKUN A.A. (MRS)\nAsst. Chief Admin & HR\nFor: Human Resource Officer")

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

# ----------------- AI EXTRACTION -----------------
genai.configure(api_key=st.secrets["GEMINI_API_KEY"])

def analyze_memo_image(uploaded_file):
    image = Image.open(uploaded_file)
    
    # THE FIX IS RIGHT HERE - using the latest model
    model = genai.GenerativeModel('gemini-1.5-flash-latest')
    
    prompt = """
    You are an expert administrative assistant parsing official government letters. 
    CRITICAL INSTRUCTION: Completely IGNORE all handwritten notes, signatures, scribbles, or minute annotations. 
    Base your extraction STRICTLY and ONLY on the formal typed/printed text of the document.
    
    Extract the following and provide the answer in the EXACT specified JSON format. No additional markdown.
    {
        "subject_line": "Formulate a dynamic and accurate subject referencing the typed meeting, seminar or retreat title. Standard title, Capital letters",
        "inviting_body": "Identify the primary inviting authority, commission, host body or organizer",
        "location": "Identify the destination city, district, or venue specified for the program",
        "duration_days": Calculate the number of calendar days of the training based on dates (return an integer, e.g., 3)"
    }
    """
    
    safety_settings = [
        {"category": "HARM_CATEGORY_HARASSMENT", "threshold": "BLOCK_NONE"},
        {"category": "HARM_CATEGORY_HATE_SPEECH", "threshold": "BLOCK_NONE"},
        {"category": "HARM_CATEGORY_SEXUALLY_EXPLICIT", "threshold": "BLOCK_NONE"},
        {"category": "HARM_CATEGORY_DANGEROUS_CONTENT", "threshold": "BLOCK_NONE"}
    ]
    
    response = model.generate_content([prompt, image], safety_settings=safety_settings)
    clean_json = response.text.strip().replace("```json", "").replace("```", "")
    return json.loads(clean_json)

# ----------------- APP PORTAL -----------------
st.subheader("📸 Upload Invitation or Memo")
uploaded_file = st.file_uploader("Take a photo or upload a document to read details automatically", type=["jpg", "png", "jpeg"])

if "extracted_data" not in st.session_state:
    st.session_state.extracted_data = None

if uploaded_file:
    if st.session_state.extracted_data is None or st.session_state.get('last_uploaded_name') != uploaded_file.name:
        with st.spinner("⏳ Analyzing document patterns..."):
            try:
                st.session_state.extracted_data = analyze_memo_image(uploaded_file)
                st.session_state.last_uploaded_name = uploaded_file.name
                st.success("🤖 Scan Successful! Extracted metadata below.")
            except Exception as e:
                st.error(f"⚠️ AI scanning failed: {str(e)}")
                st.session_state.extracted_data = {"subject_line": "", "inviting_body": "", "location": "Lagos", "duration_days": 3}

data = st.session_state.extracted_data if st.session_state.extracted_data else {"subject_line": "", "inviting_body": "", "location": "Lagos", "duration_days": 3}

st.markdown("### 🔍 AI Extraction Findings")
subject_line = st.text_input("Memo Subject Line", data.get("subject_line", ""))
inviting_body = st.text_input("Inviting Body", data.get("inviting_body", ""))

location = data.get("location", "Lagos").lower()
is_lagos = "ikeja" in location or "lagos" in location

travel_default = "Within Lagos" if is_lagos else "Road"
if "abuja" in location:
    travel_default = "Air"

opening_options = [
    f"In reference to the invitation received from {inviting_body},",
    f"On the strength of the invitation received from {inviting_body},",
    f"In view of the invitation received from {inviting_body},",
    f"Sequel to the invitation received from {inviting_body},"
]
selected_opening = st.selectbox("📝 Opening Sentence Variation", opening_options)

st.subheader("👥 Participant Card List")

if "participants_list" not in st.session_state:
    st.session_state.participants_list = [{"Name": "", "Title": "", "GL": "GL 13-14", "Travel_Mode": travel_default, "Days": data.get("duration_days", 3)}]

col_btn1, col_btn2 = st.columns(2)
with col_btn1:
    if st.button("➕ Add Officer"):
        st.session_state.participants_list.append({"Name": "", "Title": "", "GL": "GL 10-12", "Travel_Mode": travel_default, "Days": data.get("duration_days", 3)})
with col_btn2:
    if st.button("🗑️ Reset List"):
        st.session_state.participants_list = [{"Name": "", "Title": "", "GL": "GL 13-14", "Travel_Mode": travel_default, "Days": data.get("duration_days", 3)}]

for index, p in enumerate(st.session_state.participants_list):
    st.markdown(f"**Participant #{index+1}**")
    col1, col2 = st.columns(2)
    with col1:
        p["Name"] = st.text_input(f"Name", p["Name"], key=f"pname_{index}", placeholder="e.g. OSIKOYA ADESINA")
        p["Title"] = st.text_input(f"Title", p["Title"], key=f"ptitle_{index}", placeholder="e.g. Council Manager")
    with col2:
        p["GL"] = st.selectbox(f"Grade Level", ["GL 15-17", "GL 13-14", "GL 10-12", "GL 07-09"], index=["GL 15-17", "GL 13-14", "GL 10-12", "GL 07-09"].index(p["GL"]), key=f"pgl_{index}")
        p["Travel_Mode"] = st.selectbox(f"Mode", ["Within Lagos", "Air", "Road"], index=["Within Lagos", "Air", "Road"].index(p["Travel_Mode"]), key=f"ptravel_{index}")
        p["Days"] = st.number_input(f"Days", min_value=1, value=int(p["Days"]), key=f"pdays_{index}")

st.subheader("💰 Optional Special Allowances")
col_f1, col_f2 = st.columns(2)
with col_f1:
    participatory_fee = st.number_input("Participatory Fee (PER PARTICIPANT ₦)", value=0, min_value=0)
with col_f2:
    protocol_fee = st.number_input("Protocol (Total ₦)", value=0, min_value=0)

if st.button("📄 Generate & Download Document"):
    if not any(p["Name"].strip() for p in st.session_state.participants_list):
        st.error("⚠️ Add at least one Officer to calculate the allowances.")
    else:
        current_date = date.today().strftime("%dth %B, %Y")
        docx_file = generate_memo_docx(subject_line, current_date, selected_opening, st.session_state.participants_list, participatory_fee, protocol_fee)
        
        if docx_file:
            st.success("✅ Document Generation Successful!")
            st.download_button(
                label="📥 Tap here to Download memo (.docx)",
                data=docx_file,
                file_name="L&D_Draft_Memo.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            )
