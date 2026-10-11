import google.generativeai as genai
import json
from PIL import Image
import streamlit as st
import io
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from datetime import date

# --- PAGE SETUP & PASSWORD PROTECTION ---
st.set_page_config(page_title="L&D Memo App", page_icon="📝", layout="centered")

# The code asks Streamlit's hidden vault for the password
user_password = st.sidebar.text_input("Enter Office Password", type="password")

if user_password != st.secrets["APP_PASSWORD"]:
    st.warning("🔒 Please enter the correct password in the sidebar to access the L&D Memo Generator.")
    st.stop() # Stops the app from loading until the password is correct

# --- 1. ESTACODE ENGINE ---
def calculate_allowance(grade_level, days, travel_type="Air"):
    estacode_rates = {
        "GL 15-17": {"Air": {"Accom": 165000, "Airfare": 440000, "Transp": 36000, "Incidental": 27000},
                     "Road": {"Accom": 165000, "Airfare": 0, "Transp": 54000, "Incidental": 27000}},
        "GL 13-14": {"Air": {"Accom": 124875, "Airfare": 440000, "Transp": 32400, "Incidental": 20250},
                     "Road": {"Accom": 124875, "Airfare": 0, "Transp": 43200, "Incidental": 20250}},
        "GL 10-12": {"Air": {"Accom": 82800, "Airfare": 440000, "Transp": 21600, "Incidental": 15187.50},
                     "Road": {"Accom": 82800, "Airfare": 0, "Transp": 36000, "Incidental": 15187.50}},
        "GL 07-09": {"Air": {"Accom": 62100, "Airfare": 440000, "Transp": 14400, "Incidental": 10125},
                     "Road": {"Accom": 62100, "Airfare": 0, "Transp": 27000, "Incidental": 10125}}
    }
    rates = estacode_rates.get(grade_level, estacode_rates["GL 10-12"])
    selected_rates = rates[travel_type]
    
    accom_total = selected_rates["Accom"] * days
    sub_total = selected_rates["Airfare"] + accom_total + selected_rates["Transp"] + selected_rates["Incidental"]
    
    return {
        "Accom_Total": accom_total,
        "Airfare": selected_rates["Airfare"],
        "Transp": selected_rates["Transp"],
        "Incidental": selected_rates["Incidental"],
        "Sub_Total": sub_total
    }

# --- 2. DOCUMENT GENERATOR ---
def generate_memo_docx(subject, memo_date, opening_phrase, participants, part_fee, protocol):
    doc = Document()
    header = doc.add_paragraph()
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header.add_run("AGEGE LOCAL GOVERNMENT\nMEMORANDUM").bold = True
    header.runs[0].font.size = Pt(14)

    doc.add_paragraph(f"From: Learning & Development Unit\nAdmin & HR Department\nAgege Local Government\nRef: AGLG/GEN/177/007\n\nTo: The Executive Chairman\nThru: The Council Manager\nThru: The Human Resource Officer\nDate: {memo_date}")
    doc.add_paragraph().add_run(f"RE: {subject.upper()}").bold = True
    doc.add_paragraph(f"{opening_phrase} I am directed to respectfully request approval for the release of funds in respect of the above-referenced programme. The proposed expenditure is detailed below:")

    participants_subtotal = 0

    for index, p in enumerate(participants, start=1):
        # Skip empty rows if accidentally left blank
        if not p.get('Name'): continue 
        
        costs = calculate_allowance(p['GL'], p['Days'], p['Travel_Mode'])
        participants_subtotal += costs['Sub_Total']
        
        doc.add_paragraph().add_run(f"{index}. {p['Name']} | {p['Title']} [{p['GL']}]").bold = True
        
        if costs['Airfare'] > 0:
            doc.add_paragraph(f"(a) Airfare: ₦{costs['Airfare']:,.2f}")
        doc.add_paragraph(f"(b) Accommodation and feeding: ₦{costs['Accom_Total']:,.2f}")
        doc.add_paragraph(f"(c) Transportation: ₦{costs['Transp']:,.2f}")
        doc.add_paragraph(f"(d) Incidental: ₦{costs['Incidental']:,.2f}")
        doc.add_paragraph(f"SUB-TOTAL: ₦{costs['Sub_Total']:,.2f}").runs[0].bold = True

    grand_total = participants_subtotal + part_fee + protocol
    doc.add_paragraph(f"Participants Sub-total: ₦{participants_subtotal:,.2f}")
    if part_fee > 0: doc.add_paragraph(f"Participatory Fees: ₦{part_fee:,.2f}")
    if protocol > 0: doc.add_paragraph(f"Protocol: ₦{protocol:,.2f}")
    doc.add_paragraph(f"GRAND TOTAL: ₦{grand_total:,.2f}").runs[0].bold = True

    doc.add_paragraph(f"Your kind approval is hereby sought for the sum of ₦{grand_total:,.2f}.\n\nThank you.\n\nOLATOKUN A.A. (MRS)\nAsst. Chief Admin & HR\nFor: Human Resource Officer")

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

# --- 3. AI EXTRACTOR & UI ---
genai.configure(api_key=st.secrets["GEMINI_API_KEY"])

def analyze_memo_image(uploaded_file):
    image = Image.open(uploaded_file)
    model = genai.GenerativeModel('gemini-1.5-flash')
    prompt = """
    You are an expert administrative assistant. Read the attached official invitation memo.
    Extract the following details and return ONLY a valid JSON object. Do not include markdown formatting.
    {
        "subject_line": "The exact theme or subject of the workshop/meeting",
        "inviting_body": "The organization or ministry that sent the invitation",
        "location": "The city where the event will take place",
        "duration_days": The number of days the event lasts (return as an integer, e.g., 3)
    }
    """
    response = model.generate_content([prompt, image])
    clean_json = response.text.strip().replace("```json", "").replace("```", "")
    return json.loads(clean_json)

st.title("L&D Memo Generator")
st.write("Upload a memo to auto-extract details and generate the official .docx file.")

uploaded_file = st.file_uploader("Upload Invitation Memo (JPG/PNG)", type=["jpg", "png", "jpeg"])

if "extracted_data" not in st.session_state:
    st.session_state.extracted_data = None

if uploaded_file:
    if st.session_state.extracted_data is None or st.session_state.get('last_uploaded_name') != uploaded_file.name:
        with st.spinner("AI is reading the memo..."):
            try:
                st.session_state.extracted_data = analyze_memo_image(uploaded_file)
                st.session_state.last_uploaded_name = uploaded_file.name
                st.success("Details extracted successfully!")
            except Exception as e:
                st.error("Failed to read the image. Please enter details manually.")
                st.session_state.extracted_data = {"subject_line": "", "inviting_body": "", "location": "", "duration_days": 3}

    data = st.session_state.extracted_data
    subject_line = st.text_input("Memo Subject Line", data.get("subject_line", ""))
    inviting_body = st.text_input("Inviting Body", data.get("inviting_body", ""))
    
    opening_options = [
        f"In reference to the invitation received from the {inviting_body},",
        f"On the strength of the invitation received from the {inviting_body},",
        f"In view of the invitation received from the {inviting_body},",
        f"Sequel to the invitation received from the {inviting_body},"
    ]
    selected_opening = st.selectbox("Select Opening Phrase", opening_options)
    
    st.subheader("Participant List")
    auto_travel_mode = "Air" if "abuja" in data.get("location", "").lower() else "Road"
    auto_days = data.get("duration_days", 3)
    
    # Starting with a blank row instead of hardcoded names
    blank_participants = [
        {"Name": "", "Title": "", "GL": "GL 10-12", "Travel_Mode": auto_travel_mode, "Days": auto_days}
    ]
    
    edited_participants = st.data_editor(blank_participants, num_rows="dynamic")
    
    st.subheader("Additional Fees")
    col1, col2 = st.columns(2)
    with col1:
        participatory_fee = st.number_input("Participatory Fees (Total ₦)", value=0)
    with col2:
        protocol_fee = st.number_input("Protocol (Total ₦)", value=100000)

    if st.button("Generate Memo"):
        current_date = date.today().strftime("%dth %B, %Y")
        docx_file = generate_memo_docx(subject_line, current_date, selected_opening, edited_participants, participatory_fee, protocol_fee)
        
        st.download_button(
            label="📥 Download Draft Memo (.docx)",
            data=docx_file,
            file_name="LD_Draft_Memo.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )